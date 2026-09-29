import io

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from apps.auditoria.models import Accion, RegistroAuditoria
from apps.empleados import importacion
from apps.empleados.models import Empleado

ENCABEZADO = "numero_empleado,nombre,segundo_nombre,apellido_paterno,apellido_materno,ssn,fecha_empleo,tipo_pago,tarifa,departamento,aplica_choferil,banco_ruta,banco_cuenta,deposito_directo\n"


def _csv(*filas, encabezado=ENCABEZADO):
    return SimpleUploadedFile("empleados.csv", (encabezado + "\n".join(filas)).encode("utf-8"), content_type="text/csv")


@pytest.mark.django_db
def test_solo_validar_no_importa(compania, preparador):
    archivo = _csv("1,José,Luis,Santana,Rivera,123-45-6789,2015-03-01,hora,10.50,Cocina,si,,,no")
    resultado = importacion.procesar(archivo, compania, preparador, solo_validar=True)
    assert resultado.ok, resultado.errores
    assert resultado.departamentos_nuevos == {"Cocina"}
    assert Empleado.objects.count() == 0


@pytest.mark.django_db
def test_importa_y_crea_departamentos(compania, preparador):
    archivo = _csv(
        "1,José,Luis,Santana,Rivera,123-45-6789,2015-03-01,hora,10.50,Cocina,si,021502011,000123456,si",
        "2,Ana,,Pérez,,234567890,03/15/2021,salario,800,Cocina,no,,,no",
    )
    resultado = importacion.procesar(archivo, compania, preparador, solo_validar=False)
    assert resultado.ok, resultado.errores
    assert len(resultado.creados) == 2
    jose = Empleado.objects.get(numero_empleado="1")
    assert jose.ssn.revelar() == "123456789"
    assert jose.regimen_efectivo == "anterior"
    assert jose.aplica_choferil and jose.deposito_directo
    assert jose.banco_cuenta.revelar() == "000123456"
    assert jose.departamento.nombre == "Cocina"
    ana = Empleado.objects.get(numero_empleado="2")
    assert ana.regimen_efectivo == "ley4"
    assert ana.departamento == jose.departamento
    assert compania.departamentos.count() == 1


@pytest.mark.django_db
def test_errores_no_importan_nada(compania, preparador):
    archivo = _csv(
        "1,José,,Santana,,123-45-6789,2015-03-01,hora,10.50,,no,,,no",
        "2,Ana,,Pérez,,000-12-3456,2021-03-15,hora,10,,no,,,no",
        "3,Luis,,Díaz,,123-45-6789,2021-03-15,quincenal,-5,,talvez,,,no",
        "1,Rosa,,Vega,,345-67-8901,,hora,10,,no,,,si",
    )
    resultado = importacion.procesar(archivo, compania, preparador, solo_validar=False)
    assert not resultado.ok
    assert Empleado.objects.count() == 0
    columnas = {(e.fila, e.columna) for e in resultado.errores}
    assert (3, "ssn") in columnas  # SSN inválido
    assert (4, "aplica_choferil") in columnas  # valor no es si/no
    assert (4, "tipo_pago") in columnas
    assert (4, "tarifa") in columnas
    assert (5, "fecha_empleo") in columnas  # requerida
    assert (5, "banco_ruta") in columnas  # depósito directo sin ruta


@pytest.mark.django_db
def test_duplicados_dentro_del_archivo(compania, preparador):
    archivo = _csv(
        "1,José,,Santana,,123-45-6789,2015-03-01,hora,10.50,,no,,,no",
        "2,Ana,,Pérez,,123456789,2021-03-15,hora,10,,no,,,no",
    )
    resultado = importacion.procesar(archivo, compania, preparador, solo_validar=True)
    assert [(e.fila, e.columna) for e in resultado.errores] == [(3, "ssn")]


@pytest.mark.django_db
def test_faltan_columnas(compania, preparador):
    archivo = _csv("1,José", encabezado="numero_empleado,nombre\n")
    resultado = importacion.procesar(archivo, compania, preparador)
    assert not resultado.ok
    assert "ssn" in resultado.errores[0].columna


@pytest.mark.django_db
def test_excel_con_ssn_numerico_conserva_ceros(compania, preparador):
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Empleados"
    hoja.append(["Número Empleado", "Nombre", "Apellido Paterno", "SSN", "Fecha Empleo", "Tipo Pago", "Tarifa"])
    from datetime import datetime

    hoja.append(["7", "Carmen", "Ortiz", 12345678, datetime(2018, 7, 1), "Por hora", 9.5])
    salida = io.BytesIO()
    libro.save(salida)
    archivo = SimpleUploadedFile("empleados.xlsx", salida.getvalue())
    resultado = importacion.procesar(archivo, compania, preparador, solo_validar=False)
    assert resultado.ok, resultado.errores
    carmen = Empleado.objects.get(numero_empleado="7")
    assert carmen.ssn.revelar() == "012345678"
    assert str(carmen.tarifa) == "9.5000"


@pytest.mark.django_db
def test_plantilla_descargable(cliente_preparador):
    respuesta = cliente_preparador.get(reverse("empleados:plantilla"))
    assert respuesta.status_code == 200
    libro = load_workbook(io.BytesIO(respuesta.content))
    encabezado = [c.value for c in libro["Empleados"][1]]
    assert encabezado[:6] == ["numero_empleado", "nombre", "segundo_nombre", "apellido_paterno", "apellido_materno", "ssn"]
    assert "Instrucciones" in libro.sheetnames


@pytest.mark.django_db
def test_importar_por_pantalla_registra_auditoria(cliente_preparador, compania):
    archivo = _csv("1,José,,Santana,,123-45-6789,2015-03-01,hora,10.50,,no,,,no")
    respuesta = cliente_preparador.post(reverse("empleados:importar"), {"archivo": archivo})
    assert respuesta.status_code == 302
    registro = RegistroAuditoria.objects.get(accion=Accion.EMPLEADOS_IMPORTADOS)
    assert "123456789" not in str(registro.cambios)
