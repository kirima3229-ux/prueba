import io
from datetime import date, timedelta
from decimal import Decimal as D

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from apps.auditoria.models import Accion, RegistroAuditoria
from apps.companias.models import ClasificacionCFSE, TasaCFSE
from apps.nomina import servicios
from apps.nomina.models import EntradaDeduccion, EntradaIngreso, LineaResultado
from apps.parametros.models import ConceptoDeduccion, ConceptoIngreso
from apps.planillas import datos
from apps.planillas.models import Radicacion

from .conftest import crear_empleado
from .test_nomina import compania_lista, entrar_horas  # noqa: F401

T3 = (date(2026, 7, 1), date(2026, 9, 30))


@pytest.fixture
def escenario(compania_lista, preparador, admin):  # noqa: F811
    c = compania_lista
    meseros = ClasificacionCFSE.objects.create(compania=c, codigo="9079", descripcion="Restaurantes")
    oficina = ClasificacionCFSE.objects.create(compania=c, codigo="8810", descripcion="Oficina")
    TasaCFSE.objects.create(clasificacion=meseros, anio=2026, tasa_por_100=D("2.50"))
    TasaCFSE.objects.create(clasificacion=oficina, anio=2026, tasa_por_100=D("0.30"))
    ana = crear_empleado(c, numero="1", ssn="234567890", nombre="Ana", tarifa="12", clasificacion_cfse=meseros,
                         aplica_choferil=True)
    beto = crear_empleado(c, numero="2", ssn="345678901", nombre="Beto", tipo_pago="salario", tarifa="5000",
                          horas_regulares_periodo="40", clasificacion_cfse=oficina)
    propinas = ConceptoIngreso.objects.get(codigo="propinas")
    comisiones = ConceptoIngreso.objects.get(codigo="comisiones")
    reembolso = ConceptoIngreso.objects.get(codigo="reembolso")
    plan = ConceptoDeduccion.objects.get(codigo="plan_medico")
    periodos = []
    for semana in range(4):  # cuatro semanas de julio
        inicio = date(2026, 6, 29) + timedelta(weeks=semana)
        p = servicios.crear_periodo(compania=c, inicio=inicio, fin=inicio + timedelta(days=6),
                                    fecha_pago=inicio + timedelta(days=11), usuario=preparador)
        e_ana = entrar_horas(p, ana, horas_regulares="40")
        EntradaIngreso.objects.create(entrada=e_ana, concepto=propinas, monto=D("100"))
        EntradaIngreso.objects.create(entrada=e_ana, concepto=reembolso, monto=D("20"))
        e_beto = p.entradas.get(empleado=beto)
        EntradaIngreso.objects.create(entrada=e_beto, concepto=comisiones, monto=D("250"))
        EntradaDeduccion.objects.create(entrada=e_beto, concepto=plan, monto=D("50"))
        servicios.marcar_modificado(p)
        assert servicios.calcular_periodo(p, preparador) == []
        servicios.cerrar_periodo(p, preparador)
        periodos.append(p)
    servicios.reversar_periodo(periodos[3], "Prueba", admin)  # la cuarta no cuenta
    return c, ana, beto, periodos


def _suma(c, codigo, empleado=None, campo="monto"):
    q = LineaResultado.objects.filter(resultado__periodo__compania=c, codigo=codigo,
                                      resultado__periodo__fecha_pago__year=2026)
    if empleado:
        q = q.filter(resultado__empleado=empleado)
    return sum((getattr(l, campo) or D("0")) for l in q)


@pytest.mark.django_db
def test_w2pr(escenario):
    c, ana, beto, _ = escenario
    w2 = {d.empleado.pk: d for d in datos.w2pr(c, 2026)}
    a, b = w2[ana.pk], w2[beto.pk]
    assert (a.sueldos, a.propinas, a.reembolsos, a.comisiones) == (D("1440.00"), D("300.00"), D("60.00"), D("0"))
    assert (b.sueldos, b.comisiones) == (D("15000.00"), D("750.00"))
    assert b.aportaciones == {"Plan médico (Sección 125)": D("150.00")}
    assert b.sujeto_pr == D("15600.00")  # 15,750 − 150 del plan médico
    assert a.retenido_pr == _suma(c, "retencion_pr", ana) and b.retenido_pr == _suma(c, "retencion_pr", beto)
    assert a.propinas_ss == D("300.00") and a.salarios_ss + a.propinas_ss == _suma(c, "ss_empleado", ana, "base")
    assert b.salarios_ss == D("15600.00")  # 3 × (5000 + 250 − 50 del plan antes de FICA)
    assert b.retenido_ss == _suma(c, "ss_empleado", beto)
    assert b.retenido_medicare == _suma(c, "medicare_empleado", beto)


@pytest.mark.django_db
def test_r1b_r3_941_cuadran(escenario):
    c, *_ = escenario
    meses = datos.r1b(c, 2026, 3)
    assert meses[0].mes == 7 and meses[0].retenido == _suma(c, "retencion_pr")
    assert meses[0].pagado == D("1740.00") + D("15750.00")  # sin reembolsos
    p = datos.f941(c, 2026, 3)
    assert p.total_mensual == p.l10 and abs(p.l7) < D("1.00")
    assert p.empleados_12 == 0  # línea 1: el 12 de septiembre; las nóminas del escenario son de julio
    assert datos.empleados_al_12(c, 2026, 7) == 2
    assert p.salarios_ss + p.propinas_ss == _suma(c, "ss_empleado", campo="base")
    from apps.planillas.views import _t_r3

    tabla = _t_r3(c, 2026, None, False)
    assert tabla.totales[1:] == [D("0"), D("0"), D("0")]  # pagado, sujeto y retenido: trimestres = W-2PR


@pytest.mark.django_db
def test_940_y_dtrh_con_topes(escenario):
    c, ana, beto, _ = escenario
    p = datos.f940(c, 2026)
    # Beto ganó 15,600 sujetos a FUTA: 7,000 tributables y 8,600 en exceso.
    assert p.exceso_7000 == D("8600.00")
    assert p.l7 == D("7000.00") + D("1740.00") and p.l8 == p.futa_calculado
    filas, meses = datos.dtrh(c, 2026, 3)
    por = {f.empleado.pk: f for f in filas}
    assert por[beto.pk].salarios == D("15600.00") and por[beto.pk].tributable_desempleo == D("7000.00")
    assert por[beto.pk].tributable_sinot == D("9000.00")
    assert por[ana.pk].tributable_desempleo == por[ana.pk].salarios == D("1740.00")
    assert meses[0] == (7, 2)


@pytest.mark.django_db
def test_choferil_y_cfse(escenario):
    c, ana, *_ = escenario
    filas = datos.choferil(c, 2026, 3)
    assert [(f.empleado.pk, f.semanas) for f in filas] == [(ana.pk, 3)]
    assert filas[0].empleado_monto == _suma(c, "choferil_empleado")
    cfse = {f.clasificacion: f for f in datos.cfse(c, *T3)}
    assert set(cfse) == {"9079", "8810"}
    assert cfse["9079"].empleados == 1 and cfse["9079"].prima == _suma(c, "cfse", ana)
    assert cfse["9079"].tasa == D("2.5")


@pytest.mark.django_db
@pytest.mark.parametrize("codigo", ["r1b", "941", "dtrh", "choferil", "w2pr", "r3", "940", "servicios", "cfse"])
def test_pantallas(cliente_lectura, escenario, codigo):
    url = reverse("planillas:planilla", args=[codigo]) + "?anio=2026&trimestre=3"
    respuesta = cliente_lectura.get(url)
    html = respuesta.content.decode()
    assert respuesta.status_code == 200 and "234567890" not in html and "234-56-7890" not in html
    assert cliente_lectura.get(url + "&formato=pdf").content.startswith(b"%PDF")
    assert cliente_lectura.get(url + "&formato=xlsx").status_code == 200


@pytest.mark.django_db
def test_excel_para_la_agencia(cliente_lectura, cliente_preparador, escenario):
    url = reverse("planillas:planilla", args=["dtrh"]) + "?anio=2026&trimestre=3&formato=agencia"
    assert cliente_lectura.get(url).status_code == 403
    hoja = load_workbook(io.BytesIO(cliente_preparador.get(url).content)).active
    assert "234-56-7890" in [hoja.cell(row=r, column=1).value for r in range(2, hoja.max_row + 1)]
    assert RegistroAuditoria.objects.filter(accion=Accion.ARCHIVO_GENERADO,
                                            descripcion__contains="identificación completa").exists()
    assert cliente_preparador.get(reverse("planillas:planilla", args=["941"]) + "?formato=agencia").status_code == 404


@pytest.mark.django_db
def test_radicacion_e_inicio(cliente_preparador, cliente_lectura, escenario):
    url = reverse("planillas:planilla", args=["941"])
    assert cliente_lectura.post(url, {"anio": "2026", "trimestre": "3"}).status_code == 403
    respuesta = cliente_preparador.post(url, {"anio": "2026", "trimestre": "3", "fecha": "10/20/2026",
                                              "confirmacion": "EFTPS-123", "monto": "1500"})
    assert respuesta.status_code == 302
    r = Radicacion.objects.get()
    assert (r.tipo, r.anio, r.trimestre, r.confirmacion) == ("941", 2026, 3, "EFTPS-123")
    assert RegistroAuditoria.objects.filter(accion=Accion.PLANILLA_RADICADA).exists()
    html = cliente_lectura.get(reverse("planillas:inicio") + "?anio=2026").content.decode()
    assert "Radicada" in html and "Sin nóminas" in html  # T1 y T2 de 2026 no tienen nóminas
    assert cliente_lectura.get(reverse("planillas:planilla", args=["nada"])).status_code == 404
