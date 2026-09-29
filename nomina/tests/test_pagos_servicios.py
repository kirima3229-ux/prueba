import io
from datetime import date, timedelta
from decimal import Decimal as D

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import load_workbook

from apps.auditoria.models import Accion, RegistroAuditoria
from apps.servicios import importacion, pagos
from apps.servicios.models import ConfigRetencionServicios, ErrorPagoInmutable, PagoServicio, ProveedorServicios

from .test_servicios import crear_proveedor

HOY = date.today()


def pagar(proveedor, monto, usuario=None, fecha=HOY, **extra):
    pago, _avisos = pagos.registrar(proveedor=proveedor, fecha=fecha, monto=D(monto), usuario=usuario,
                                    numero_cheque="100", **extra)
    return pago


@pytest.mark.django_db
def test_configuracion_inicial_cargada_por_verificar():
    for anio in (2025, 2026):
        c = ConfigRetencionServicios.objects.get(anio=anio)
        assert c.tasa_general == D("10.00") and c.exencion_anual == D("500.00")
        assert c.estado == "por_verificar"


@pytest.mark.django_db
def test_pagos_sucesivos_consumen_la_exencion(compania, preparador):
    proveedor = crear_proveedor(compania)
    p1 = pagar(proveedor, "300", preparador)
    p2 = pagar(proveedor, "300", preparador)
    p3 = pagar(proveedor, "1000", preparador)
    assert (p1.retencion, p2.retencion, p3.retencion) == (D("0.00"), D("10.00"), D("100.00"))
    assert p2.acumulado_previo == D("300.00") and p2.exencion_aplicada == D("200.00")
    assert not p1.config_verificada


@pytest.mark.django_db
def test_exencion_es_por_proveedor_y_por_anio(compania, preparador):
    a = crear_proveedor(compania)
    b = crear_proveedor(compania, numero="P-2", identificacion="234567890")
    pagar(a, "500", preparador)
    assert pagar(b, "500", preparador).retencion == D("0.00")
    anterior = ConfigRetencionServicios.objects.get(anio=2025)
    assert pagar(a, "500", preparador, fecha=date(anterior.anio, 6, 1)).retencion == D("0.00")


@pytest.mark.django_db
def test_relevo_segun_fecha_del_pago(compania, preparador):
    proveedor = crear_proveedor(
        compania, relevo="parcial", relevo_porcentaje="3.00", relevo_vigente_hasta=HOY - timedelta(days=10)
    )
    antes = pagar(proveedor, "1500", preparador, fecha=HOY - timedelta(days=20))
    despues = pagar(proveedor, "1000", preparador, fecha=HOY)
    assert antes.tasa_aplicada == D("3.00") and antes.retencion == D("30.00")
    assert despues.tasa_aplicada == D("10.00") and despues.retencion == D("100.00")


@pytest.mark.django_db
def test_declaracion_jurada_sin_retencion(compania, preparador):
    proveedor = crear_proveedor(compania, relevo="declaracion_jurada")
    pago = pagar(proveedor, "5000", preparador)
    assert pago.retencion == D("0.00") and pago.tratamiento == "declaracion_jurada"


@pytest.mark.django_db
def test_sin_configuracion_del_anio(compania, preparador):
    proveedor = crear_proveedor(compania)
    with pytest.raises(pagos.ErrorPago, match="No hay configuración"):
        pagar(proveedor, "100", preparador, fecha=date(2019, 1, 1))


@pytest.mark.django_db
def test_pago_exento_requiere_motivo(compania, preparador):
    proveedor = crear_proveedor(compania)
    with pytest.raises(pagos.ErrorPago):
        pagar(proveedor, "900", preparador, exento=True)
    pago = pagar(proveedor, "900", preparador, exento=True, motivo_exencion="Servicio 1062.03(b)")
    assert pago.retencion == D("0.00")


@pytest.mark.django_db
def test_pago_inmutable(compania, preparador):
    pago = pagar(crear_proveedor(compania), "800", preparador)
    pago.monto = D("1")
    with pytest.raises(ErrorPagoInmutable):
        pago.save()
    with pytest.raises(ErrorPagoInmutable):
        pago.delete()


@pytest.mark.django_db
def test_registrar_por_pantalla_con_vista_previa(cliente_preparador, compania):
    proveedor = crear_proveedor(compania)
    datos = {"proveedor": proveedor.pk, "fecha": HOY.isoformat(), "monto": "1500", "metodo": "cheque", "numero_cheque": "501"}
    vista = cliente_preparador.post(reverse("servicios:pago_nuevo"), datos, HTTP_HX_REQUEST="true")
    html = vista.content.decode()
    assert "$100.00" in html and "<html" not in html
    assert PagoServicio.objects.count() == 0  # la vista previa no guarda

    respuesta = cliente_preparador.post(reverse("servicios:pago_nuevo"), {**datos, "confirmar": "1"})
    pago = PagoServicio.objects.get()
    assert respuesta["Location"] == reverse("servicios:pago_detalle", args=[pago.pk])
    assert pago.retencion == D("100.00") and pago.neto == D("1400.00")
    assert RegistroAuditoria.objects.filter(accion=Accion.PAGO_SERVICIO_REGISTRADO).exists()


@pytest.mark.django_db
def test_cheque_requiere_numero(cliente_preparador, compania):
    proveedor = crear_proveedor(compania)
    datos = {"proveedor": proveedor.pk, "fecha": HOY.isoformat(), "monto": "100", "metodo": "cheque", "confirmar": "1"}
    assert "Indique el número de cheque" in cliente_preparador.post(reverse("servicios:pago_nuevo"), datos).content.decode()


@pytest.mark.django_db
def test_anular_y_conciliacion(cliente_preparador, compania, preparador):
    proveedor = crear_proveedor(compania)
    p1 = pagar(proveedor, "500", preparador)
    p2 = pagar(proveedor, "1000", preparador)  # retiene 100
    assert p2.retencion == D("100.00")
    cliente_preparador.post(reverse("servicios:pago_anular", args=[p1.pk]), {"motivo": "Factura duplicada"})
    p1.refresh_from_db()
    assert p1.estado == "anulado" and p1.anulado_por == preparador
    assert RegistroAuditoria.objects.filter(accion=Accion.PAGO_SERVICIO_ANULADO).exists()
    # Sin p1, p2 debió retener solo $50: el resumen lo marca para revisión.
    html = cliente_preparador.get(reverse("servicios:resumen")).content.decode()
    assert "Revisar" in html
    # El siguiente pago ya no cuenta el anulado.
    assert pagar(proveedor, "100", preparador).acumulado_previo == D("1000.00")


@pytest.mark.django_db
def test_solo_lectura_no_registra_ni_anula(cliente_lectura, compania, preparador):
    pago = pagar(crear_proveedor(compania), "800", preparador)
    assert cliente_lectura.get(reverse("servicios:pagos")).status_code == 200
    assert cliente_lectura.get(reverse("servicios:pago_detalle", args=[pago.pk])).status_code == 200
    assert cliente_lectura.get(reverse("servicios:pago_nuevo")).status_code == 403
    assert cliente_lectura.post(reverse("servicios:pago_anular", args=[pago.pk]), {"motivo": "x"}).status_code == 403


@pytest.mark.django_db
def test_aislamiento_por_compania(cliente_preparador, compania, otra_compania, preparador):
    ajeno = crear_proveedor(otra_compania, numero="X-1")
    pago_ajeno = pagar(ajeno, "800", preparador)
    assert cliente_preparador.get(reverse("servicios:pago_detalle", args=[pago_ajeno.pk])).status_code == 404
    datos = {"proveedor": ajeno.pk, "fecha": HOY.isoformat(), "monto": "100", "metodo": "otro", "confirmar": "1"}
    respuesta = cliente_preparador.post(reverse("servicios:pago_nuevo"), datos)
    assert respuesta.status_code == 200 and PagoServicio.objects.count() == 1


@pytest.mark.django_db
def test_configuracion_solo_admin(cliente_admin, cliente_preparador, admin):
    config = ConfigRetencionServicios.objects.get(anio=2026)
    assert cliente_preparador.get(reverse("servicios:config_lista")).status_code == 403
    cliente_admin.post(
        reverse("servicios:config_editar", args=[config.pk]),
        {"anio": 2026, "tasa_general": "10.00", "exencion_anual": "500.00", "marcar_verificado": "on"},
    )
    config.refresh_from_db()
    assert config.estado == "verificado" and config.verificado_por == admin
    assert RegistroAuditoria.objects.filter(accion=Accion.CONFIGURACION_MODIFICADA).exists()


@pytest.mark.django_db
def test_resumen_excel_protege_contra_formulas(cliente_preparador, compania, preparador):
    proveedor = crear_proveedor(compania, nombre="=HYPERLINK(\"http://x\")", apellido_paterno="Malicioso")
    pagar(proveedor, "1500", preparador)
    respuesta = cliente_preparador.get(reverse("servicios:resumen"), {"formato": "xlsx"})
    hoja = load_workbook(io.BytesIO(respuesta.content)).active
    assert hoja["B2"].value.startswith("'=")
    assert hoja["I2"].value == 100  # retenido
    assert hoja["D2"].value == "6789"  # solo últimos 4 del SSN
    assert RegistroAuditoria.objects.filter(accion=Accion.ARCHIVO_GENERADO).exists()


# --- Importación ------------------------------------------------------------------


def _csv(texto, nombre="archivo.csv"):
    return SimpleUploadedFile(nombre, texto.encode("utf-8"), content_type="text/csv")


@pytest.mark.django_db
def test_importar_proveedores(compania, preparador):
    archivo = _csv(
        "numero,tipo_persona,nombre,apellido_paterno,tipo_identificacion,identificacion,relevo,relevo_porcentaje,relevo_vigente_hasta\n"
        "P-1,individuo,Luis,Méndez,ssn,123-45-6789,ninguno,,\n"
        "P-2,Entidad,ACME LLC,,EIN,66-1234567,parcial,5,2030-12-31\n"
    )
    resultado = importacion.procesar_proveedores(archivo, compania, preparador, solo_validar=False)
    assert resultado.ok, resultado.errores
    acme = ProveedorServicios.objects.get(numero="P-2")
    assert acme.tipo_persona == "entidad" and acme.identificacion.revelar() == "661234567"
    assert acme.relevo_porcentaje == D("5.00")


@pytest.mark.django_db
def test_importar_proveedores_con_errores_no_guarda(compania, preparador):
    archivo = _csv(
        "numero,tipo_persona,nombre,apellido_paterno,tipo_identificacion,identificacion\n"
        "P-1,individuo,Luis,Méndez,ssn,123-45-6789\n"
        "P-2,individuo,Ana,,ssn,000-12-3456\n"
        "P-3,individuo,Rosa,Vega,ssn,123456789\n"
    )
    resultado = importacion.procesar_proveedores(archivo, compania, preparador, solo_validar=False)
    assert not resultado.ok and ProveedorServicios.objects.count() == 0
    columnas = {(e.fila, e.columna) for e in resultado.errores}
    assert (3, "apellido_paterno") in columnas and (3, "identificacion") in columnas
    assert (4, "identificacion") in columnas  # repetida en el archivo


@pytest.mark.django_db
def test_importar_pagos_en_orden_del_archivo(compania, preparador):
    crear_proveedor(compania)
    archivo = _csv(
        "numero_proveedor,fecha,monto,referencia,metodo_pago,numero_cheque\n"
        f"P-1,{HOY.isoformat()},300,F-1,cheque,1001\n"
        f"P-1,{HOY.isoformat()},\"$1,200.00\",F-2,transferencia,\n"
    )
    vista = importacion.procesar_pagos(archivo, compania, preparador, solo_validar=True)
    assert vista.ok, vista.errores
    assert [v["r"].retencion for v in vista.vista_previa] == [D("0.00"), D("100.00")]
    assert PagoServicio.objects.count() == 0

    archivo.seek(0)
    resultado = importacion.procesar_pagos(archivo, compania, preparador, solo_validar=False)
    assert [p.retencion for p in resultado.creados] == [D("0.00"), D("100.00")]
    assert all(p.origen == "importado" for p in resultado.creados)


@pytest.mark.django_db
def test_importar_pagos_errores(compania, preparador):
    crear_proveedor(compania)
    archivo = _csv(
        "numero_proveedor,fecha,monto,metodo_pago,numero_cheque,exento\n"
        f"NO-EXISTE,{HOY.isoformat()},100,otro,,no\n"
        f"P-1,{HOY.isoformat()},100,cheque,,no\n"
        f"P-1,2019-01-01,100,otro,,no\n"
        f"P-1,{HOY.isoformat()},100,otro,,si\n"
    )
    resultado = importacion.procesar_pagos(archivo, compania, preparador, solo_validar=False)
    assert PagoServicio.objects.count() == 0
    columnas = {(e.fila, e.columna) for e in resultado.errores}
    assert (2, "numero_proveedor") in columnas
    assert (3, "numero_cheque") in columnas
    assert (4, "fecha") in columnas  # año sin configuración
    assert (5, "motivo_exencion") in columnas


@pytest.mark.django_db
def test_plantillas_descargables(cliente_preparador):
    for nombre, hoja in (("servicios:plantilla_proveedores", "Proveedores"), ("servicios:plantilla_pagos", "Pagos")):
        respuesta = cliente_preparador.get(reverse(nombre))
        assert respuesta.status_code == 200
        assert hoja in load_workbook(io.BytesIO(respuesta.content)).sheetnames
