import io
from datetime import date
from decimal import Decimal as D

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from apps.auditoria.models import Accion, RegistroAuditoria
from apps.nomina.models import DeduccionRecurrente, PeriodoNomina
from apps.parametros.models import ConceptoDeduccion, ConceptoIngreso

from .conftest import crear_empleado
from .test_nomina import compania_lista  # noqa: F401  (fixture)


@pytest.fixture
def empleados(compania_lista):  # noqa: F811
    return (crear_empleado(compania_lista, numero="1", ssn="234567890", nombre="Ana", tarifa="12"),
            crear_empleado(compania_lista, numero="2", ssn="345678901", nombre="Luis", tipo_pago="salario",
                           tarifa="800", horas_regulares_periodo="40"))


def crear_por_pantalla(cliente):
    respuesta = cliente.post(reverse("nomina:nuevo"), {
        "fecha_inicio": "2026-09-07", "fecha_fin": "2026-09-13", "fecha_pago": "2026-09-18", "tipo": "regular"})
    periodo = PeriodoNomina.objects.latest("id")
    assert respuesta["Location"] == reverse("nomina:detalle", args=[periodo.pk])
    return periodo


@pytest.mark.django_db
def test_flujo_completo_por_pantalla(cliente_preparador, compania_lista, empleados):  # noqa: F811
    ana, luis = empleados
    periodo = crear_por_pantalla(cliente_preparador)
    ea, el = periodo.entradas.get(empleado=ana), periodo.entradas.get(empleado=luis)
    respuesta = cliente_preparador.post(reverse("nomina:detalle", args=[periodo.pk]), {
        "accion": "guardar",
        f"incluir_{ea.pk}": "on", f"horas_regulares_{ea.pk}": "40", f"horas_extra_semanales_{ea.pk}": "2",
        f"propinas_{ea.pk}": "50",
        f"incluir_{el.pk}": "on",
    })
    assert respuesta.status_code == 302
    ea.refresh_from_db()
    assert ea.horas_regulares == D("40") and ea.ingresos.get().monto == D("50")

    cliente_preparador.post(reverse("nomina:calcular", args=[periodo.pk]))
    periodo.refresh_from_db()
    assert periodo.estado == "calculada" and periodo.resultados.count() == 2
    ra = periodo.resultados.get(empleado=ana)
    assert ra.bruto == D("566.00")  # 480 + 2 × 12 × 1.5 + 50
    assert periodo.resultados.get(empleado=luis).bruto == D("800.00")

    html = cliente_preparador.get(reverse("nomina:detalle", args=[periodo.pk])).content.decode()
    assert "Pre-nómina" in html and "566.00" in html and "234567890" not in html

    pdf = cliente_preparador.get(reverse("nomina:talonarios", args=[periodo.pk]))
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    excel = cliente_preparador.get(reverse("nomina:registro", args=[periodo.pk]))
    hoja = load_workbook(io.BytesIO(excel.content)).active
    assert hoja["B2"].value in ("Ana Del Pueblo", "Luis Del Pueblo")
    assert hoja.cell(row=hoja.max_row, column=2).value == "TOTAL"

    cliente_preparador.post(reverse("nomina:cerrar", args=[periodo.pk]))
    periodo.refresh_from_db()
    assert periodo.estado == "cerrada"
    assert RegistroAuditoria.objects.filter(accion=Accion.NOMINA_PROCESADA).exists()
    # Cerrada: guardar horas no cambia nada.
    cliente_preparador.post(reverse("nomina:detalle", args=[periodo.pk]), {"accion": "guardar", f"horas_regulares_{ea.pk}": "1"})
    ea.refresh_from_db()
    assert ea.horas_regulares == D("40")


@pytest.mark.django_db
def test_excluir_empleado(cliente_preparador, compania_lista, empleados):  # noqa: F811
    ana, luis = empleados
    periodo = crear_por_pantalla(cliente_preparador)
    ea, el = periodo.entradas.get(empleado=ana), periodo.entradas.get(empleado=luis)
    cliente_preparador.post(reverse("nomina:detalle", args=[periodo.pk]), {
        "accion": "guardar", f"incluir_{ea.pk}": "on", f"horas_regulares_{ea.pk}": "40"})
    cliente_preparador.post(reverse("nomina:calcular", args=[periodo.pk]))
    assert list(periodo.resultados.values_list("empleado", flat=True)) == [ana.pk]


@pytest.mark.django_db
def test_valor_invalido_no_guarda_nada(cliente_preparador, compania_lista, empleados):  # noqa: F811
    ana, _ = empleados
    periodo = crear_por_pantalla(cliente_preparador)
    ea = periodo.entradas.get(empleado=ana)
    respuesta = cliente_preparador.post(reverse("nomina:detalle", args=[periodo.pk]), {
        "accion": "guardar", f"incluir_{ea.pk}": "on", f"horas_regulares_{ea.pk}": "abc"})
    assert "no es un número válido" in respuesta.content.decode()


@pytest.mark.django_db
def test_detalle_de_empleado_ingresos_y_deducciones(cliente_preparador, compania_lista, empleados):  # noqa: F811
    ana, _ = empleados
    periodo = crear_por_pantalla(cliente_preparador)
    ea = periodo.entradas.get(empleado=ana)
    url = reverse("nomina:entrada", args=[periodo.pk, ea.pk])
    reembolso = ConceptoIngreso.objects.get(codigo="reembolso")
    prestamo = ConceptoDeduccion.objects.get(codigo="prestamo")
    cliente_preparador.post(url, {"accion": "ingreso", "ing-concepto": reembolso.pk, "ing-monto": "25"})
    cliente_preparador.post(url, {"accion": "deduccion", "ded-concepto": prestamo.pk, "ded-monto": "40"})
    cliente_preparador.post(url, {"accion": "horas", "incluir": "on", "horas_regulares": "40",
                                  **{c: "0" for c in ("horas_extra_diarias", "horas_extra_semanales", "horas_septimo_dia",
                                                      "horas_periodo_alimentos", "horas_vacaciones", "horas_enfermedad")}})
    cliente_preparador.post(reverse("nomina:calcular", args=[periodo.pk]))
    r = periodo.resultados.get(empleado=ana)
    assert r.bruto == D("505.00") and r.trib_pr == D("480.00") and r.total_deducciones == D("40.00")
    assert "Reembolso" in cliente_preparador.get(url).content.decode()


@pytest.mark.django_db
def test_permisos(cliente_lectura, cliente_preparador, cliente_admin, compania_lista, empleados):  # noqa: F811
    ana, _ = empleados
    periodo = crear_por_pantalla(cliente_preparador)
    ea = periodo.entradas.get(empleado=ana)
    cliente_preparador.post(reverse("nomina:detalle", args=[periodo.pk]), {
        "accion": "guardar", f"incluir_{ea.pk}": "on", f"horas_regulares_{ea.pk}": "40"})
    assert cliente_lectura.get(reverse("nomina:lista")).status_code == 200
    assert cliente_lectura.get(reverse("nomina:detalle", args=[periodo.pk])).status_code == 200
    assert cliente_lectura.get(reverse("nomina:nuevo")).status_code == 403
    assert cliente_lectura.post(reverse("nomina:calcular", args=[periodo.pk])).status_code == 403
    cliente_lectura.post(reverse("nomina:detalle", args=[periodo.pk]), {"accion": "guardar", f"horas_regulares_{ea.pk}": "1"})
    ea.refresh_from_db()
    assert ea.horas_regulares == D("40")
    cliente_preparador.post(reverse("nomina:calcular", args=[periodo.pk]))
    cliente_preparador.post(reverse("nomina:cerrar", args=[periodo.pk]))
    # Solo el administrador reversa.
    assert cliente_preparador.post(reverse("nomina:reversar", args=[periodo.pk]), {"motivo": "x"}).status_code == 403
    respuesta = cliente_admin.post(reverse("nomina:reversar", args=[periodo.pk]), {"motivo": "Error de horas"})
    periodo.refresh_from_db()
    assert periodo.estado == "reversada" and respuesta.status_code == 302
    assert RegistroAuditoria.objects.filter(accion=Accion.NOMINA_REVERSADA).exists()


@pytest.mark.django_db
def test_otra_compania_no_accesible(cliente_preparador, otra_compania, admin):
    from apps.nomina import servicios

    ajeno = servicios.crear_periodo(compania=otra_compania, inicio=date(2026, 9, 7), fin=date(2026, 9, 13),
                                    fecha_pago=date(2026, 9, 18), usuario=admin)
    assert cliente_preparador.get(reverse("nomina:detalle", args=[ajeno.pk])).status_code == 404
    assert cliente_preparador.get(reverse("nomina:talonarios", args=[ajeno.pk])).status_code == 404


@pytest.mark.django_db
def test_deducciones_recurrentes(cliente_preparador, compania_lista, empleados):  # noqa: F811
    ana, _ = empleados
    plan = ConceptoDeduccion.objects.get(codigo="plan_medico")
    url = reverse("nomina:deducciones_empleado", args=[ana.pk])
    cliente_preparador.post(url, {"accion": "nueva", "concepto": plan.pk, "monto": "45"})
    ded = DeduccionRecurrente.objects.get(empleado=ana)
    periodo = crear_por_pantalla(cliente_preparador)
    assert periodo.entradas.get(empleado=ana).deducciones.get().monto == D("45.00")
    cliente_preparador.post(url, {"accion": "desactivar", "id": ded.pk})
    ded.refresh_from_db()
    assert not ded.activo
    assert RegistroAuditoria.objects.filter(accion=Accion.DEDUCCION_RECURRENTE).count() == 2
