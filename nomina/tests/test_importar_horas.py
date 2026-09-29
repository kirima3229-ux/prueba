import io
from datetime import date
from decimal import Decimal as D

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from apps.licencias.models import BonoNavidad  # noqa: F401
from apps.nomina import importar_horas, servicios
from apps.nomina.models import EntradaIngreso

from .conftest import crear_empleado
from .test_nomina import compania_lista, entrar_horas, periodo_semana  # noqa: F401


def _xlsx(filas, nombre="horas.xlsx"):
    libro = Workbook()
    hoja = libro.active
    for f in filas:
        hoja.append(f)
    salida = io.BytesIO()
    libro.save(salida)
    return SimpleUploadedFile(nombre, salida.getvalue())


@pytest.fixture
def periodo(compania_lista, preparador):  # noqa: F811
    crear_empleado(compania_lista, numero="1", ssn="234567890", nombre="Ana", tarifa="12")
    crear_empleado(compania_lista, numero="2", ssn="345678901", nombre="Beto", tarifa="15")
    return periodo_semana(compania_lista, preparador)


@pytest.mark.django_db
def test_plantilla_trae_empleados(periodo):
    hoja = load_workbook(io.BytesIO(importar_horas.plantilla(periodo)))["Horas"]
    assert [hoja.cell(row=r, column=1).value for r in (2, 3)] == ["1", "2"]
    assert hoja.cell(row=1, column=3).value == "horas_regulares"


@pytest.mark.django_db
def test_validar_y_aplicar(periodo):
    archivo = _xlsx([["numero_empleado", "horas_regulares", "horas_extra_semanales", "propinas", "otra"],
                     ["1", 40, 2.5, "$120.50", "x"], ["2", 32, None, None, None]])
    r = importar_horas.procesar(archivo, periodo, solo_validar=True)
    assert not r.errores and r.total_filas == 2 and r.aplicadas == 0 and r.columnas_desconocidas == ["otra"]
    assert periodo.entradas.get(empleado__numero_empleado="1").horas_regulares == 0
    archivo.seek(0)
    r = importar_horas.procesar(archivo, periodo, solo_validar=False)
    assert r.aplicadas == 2
    ana = periodo.entradas.get(empleado__numero_empleado="1")
    assert (ana.horas_regulares, ana.horas_extra_semanales) == (D("40"), D("2.5"))
    assert ana.ingresos.get(concepto__codigo="propinas").monto == D("120.50")
    beto = periodo.entradas.get(empleado__numero_empleado="2")
    assert beto.horas_regulares == D("32") and not beto.ingresos.exists()


@pytest.mark.django_db
def test_errores_no_cambian_nada(periodo):
    archivo = _xlsx([["numero_empleado", "horas_regulares"], ["1", 40], ["9", 10], ["1", 5], ["2", "abc"]])
    r = importar_horas.procesar(archivo, periodo, solo_validar=False)
    mensajes = " ".join(e.mensaje for e in r.errores)
    assert "«9» no está" in mensajes and "se repite" in mensajes and "no es una cantidad" in mensajes
    assert r.aplicadas == 0 and periodo.entradas.get(empleado__numero_empleado="1").horas_regulares == 0


@pytest.mark.django_db
def test_csv_y_columnas_ausentes_no_se_tocan(periodo):
    ana = periodo.entradas.get(empleado__numero_empleado="1")
    ana.horas_vacaciones = D("8")
    ana.save()
    csv = SimpleUploadedFile("horas.csv", "Número empleado;Horas regulares\n1;38\n".encode("utf-8"))
    r = importar_horas.procesar(csv, periodo, solo_validar=False)
    assert not r.errores
    ana.refresh_from_db()
    assert ana.horas_regulares == D("38") and ana.horas_vacaciones == D("8")


@pytest.mark.django_db
def test_pantalla(cliente_preparador, cliente_lectura, periodo):
    url = reverse("nomina:importar_horas", args=[periodo.pk])
    assert cliente_lectura.get(url).status_code == 403
    assert cliente_preparador.get(url + "?plantilla=1")["Content-Type"].startswith("application/vnd.openxml")
    archivo = _xlsx([["numero_empleado", "horas_regulares"], ["2", 40]])
    respuesta = cliente_preparador.post(url, {"archivo": archivo})
    assert respuesta.status_code == 302
    assert periodo.entradas.get(empleado__numero_empleado="2").horas_regulares == D("40")
    periodo.refresh_from_db()
    assert periodo.requiere_recalculo


# --- Datos de la nómina para licencias y bono; períodos especiales -------------------


@pytest.mark.django_db
def test_horas_y_bono_desde_nomina(compania_lista, preparador, admin):  # noqa: F811
    ana = crear_empleado(compania_lista, numero="1", ssn="234567890", nombre="Ana", tarifa="12")
    luis = crear_empleado(compania_lista, numero="2", ssn="345678901", nombre="Luis", tipo_pago="salario",
                          tarifa="800", horas_regulares_periodo="40")
    for inicio in (date(2026, 9, 7), date(2026, 9, 14), date(2026, 9, 28)):
        p = periodo_semana(compania_lista, preparador, inicio)
        entrar_horas(p, ana, horas_regulares="40")
        entrar_horas(p, luis, horas_vacaciones="8")
        p.entradas.filter(empleado=ana).first().ingresos.create(
            concepto_id=EntradaIngreso._meta.get_field("concepto").related_model.objects.get(codigo="propinas").pk,
            monto=D("100"))
        servicios.calcular_periodo(p, preparador)
        servicios.cerrar_periodo(p, preparador)
    # El de la semana del 14 se reversa: no cuenta.
    servicios.reversar_periodo(p.__class__.objects.get(fecha_inicio=date(2026, 9, 14)), "error", admin)
    horas = servicios.horas_del_mes(compania_lista, 2026, 9)
    assert horas[ana.pk] == D("40.00")  # el período del 28 termina en octubre
    assert horas[luis.pk] == D("32.00")  # asalariado: 40 regulares − 8 de vacaciones
    datos = servicios.datos_bono(compania_lista, date(2025, 10, 1), date(2026, 10, 31))
    assert datos[ana.pk] == (D("80.00"), D("960.00"))  # sin propinas
    assert datos[luis.pk] == (D("64.00"), D("1600.00"))


@pytest.mark.django_db
def test_pantallas_traen_de_la_nomina(cliente_preparador, compania_lista, preparador):  # noqa: F811
    ana = crear_empleado(compania_lista, numero="1", ssn="234567890", nombre="Ana", tarifa="12")
    p = periodo_semana(compania_lista, preparador)
    entrar_horas(p, ana, horas_regulares="37.5")
    servicios.calcular_periodo(p, preparador)
    servicios.cerrar_periodo(p, preparador)
    html = cliente_preparador.get(reverse("licencias:acumular") + "?anio=2026&mes=9&de_nomina=1").content.decode()
    assert 'value="37.50"' in html
    html = cliente_preparador.get(reverse("licencias:bono") + "?anio=2026&de_nomina=1").content.decode()
    assert 'value="37.50"' in html and 'value="450.00"' in html


@pytest.mark.django_db
def test_periodo_especial_no_paga_salario_ni_choferil(compania_lista, preparador):  # noqa: F811
    luis = crear_empleado(compania_lista, numero="2", ssn="345678901", nombre="Luis", tipo_pago="salario",
                          tarifa="800", horas_regulares_periodo="40", aplica_choferil=True)
    p = servicios.crear_periodo(compania=compania_lista, inicio=date(2026, 12, 1), fin=date(2026, 12, 15),
                                fecha_pago=date(2026, 12, 15), usuario=preparador, tipo="especial")
    entrada = p.entradas.get(empleado=luis)
    from apps.parametros.models import ConceptoIngreso

    entrada.ingresos.create(concepto=ConceptoIngreso.objects.get(codigo="bono_navidad"), monto=D("600"))
    assert servicios.calcular_periodo(p, preparador) == []
    r = p.resultados.get()
    assert r.bruto == D("600.00") and r.horas_trabajadas == 0
    assert not r.lineas.filter(codigo__startswith="choferil").exists()
