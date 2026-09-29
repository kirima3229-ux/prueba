import io
from datetime import date
from decimal import Decimal as D

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from apps.companias.models import Departamento
from apps.nomina import reportes, servicios

from .conftest import crear_empleado
from .test_nomina import compania_lista, entrar_horas, periodo_semana  # noqa: F401


@pytest.fixture
def dos_nominas(compania_lista, preparador, admin):  # noqa: F811
    cocina = Departamento.objects.create(compania=compania_lista, nombre="Cocina")
    ana = crear_empleado(compania_lista, numero="1", ssn="234567890", nombre="Ana", tarifa="12", departamento=cocina)
    beto = crear_empleado(compania_lista, numero="2", ssn="345678901", nombre="Beto", tarifa="15")
    periodos = []
    for inicio in (date(2026, 9, 7), date(2026, 9, 14), date(2026, 9, 21)):
        p = periodo_semana(compania_lista, preparador, inicio)
        entrar_horas(p, ana, horas_regulares="40")
        entrar_horas(p, beto, horas_regulares="40")
        servicios.calcular_periodo(p, preparador)
        servicios.cerrar_periodo(p, preparador)
        periodos.append(p)
    # La tercera se reversa: no debe contar.
    servicios.reversar_periodo(periodos[2], "Prueba", admin)
    # Una en borrador tampoco cuenta.
    borrador = periodo_semana(compania_lista, preparador, date(2026, 9, 28))
    entrar_horas(borrador, ana, horas_regulares="40")
    return compania_lista, ana, beto


SEP = (date(2026, 9, 1), date(2026, 10, 31))


@pytest.mark.django_db
def test_resumen(dos_nominas):
    compania, *_ = dos_nominas
    r = reportes.resumen(compania, *SEP)
    assert r.totales["ingreso"] == D("2160.00")  # 2 × (480 + 600)
    ss = next(f for f in r.grupos["patronal"] if f[0] == "ss_patrono")
    assert ss[2] == D("133.92") and ss[3] == D("6.20")
    assert r.empleados == 2 and r.nominas == 2  # la reversada y su reverso no cuentan


@pytest.mark.django_db
def test_por_empleado_y_costo(dos_nominas):
    compania, ana, beto = dos_nominas
    filas = {f["empleado_nombre"][:4]: f for f in reportes.por_empleado(compania, *SEP)}
    assert filas["Ana "]["bruto"] == D("960.00") and filas["Ana "]["horas"] == D("80.00")
    assert filas["Ana "]["columnas"][1] == D("59.52")  # SS 6.2% de 960
    assert filas["Ana "]["neto"] == filas["Ana "]["bruto"] - sum(filas["Ana "]["columnas"]) - filas["Ana "]["deducciones"]
    conceptos, filas_costo, total = reportes.costo_patronal(compania, *SEP)
    assert "Seguro Social patronal" in conceptos
    assert total["bruto"] == D("2160.00") and total["costo"] == total["bruto"] + total["patronal"]
    assert sum(f["pct_del_total"] for f in filas_costo) == D("100.00")
    _, por_depto, _ = reportes.costo_patronal(compania, *SEP, agrupar="departamento")
    assert sorted(f["nombre"] for f in por_depto) == ["Cocina", "Sin departamento"]


@pytest.mark.django_db
def test_impuestos_y_vencimientos(dos_nominas):
    compania, *_ = dos_nominas
    obligaciones = reportes.impuestos(compania, *SEP)
    irs = [o for o in obligaciones if o.agencia == "IRS" and "Seguro Social" in o.concepto]
    assert len(irs) == 1 and irs[0].vence == date(2026, 10, 15)  # mensual
    assert irs[0].monto == D("330.48")  # (6.2 + 6.2 + 1.45 + 1.45)% de 2160
    dtrh = [o for o in obligaciones if o.concepto.startswith("Desempleo estatal")]
    assert dtrh[0].vence == date(2026, 11, 2) and dtrh[0].periodo == "Trimestre 3/2026"
    compania.frecuencia_deposito_federal = "bisemanal"
    compania.save()
    irs = [o for o in reportes.impuestos(compania, *SEP) if o.agencia == "IRS" and "Seguro Social" in o.concepto]
    assert [(o.periodo, o.vence) for o in irs] == [
        ("Pago del 09/18/2026", date(2026, 9, 23)), ("Pago del 09/25/2026", date(2026, 9, 30))]


@pytest.mark.django_db
@pytest.mark.parametrize("reporte", ["resumen", "empleados", "costo", "impuestos"])
def test_pantalla_excel_pdf(cliente_lectura, dos_nominas, reporte):
    url = reverse("nomina:reportes") + f"?reporte={reporte}&rango=trimestre&anio=2026&trimestre=3"
    html = cliente_lectura.get(url).content.decode()
    assert "trimestre 3 de 2026" in html and "234567890" not in html
    excel = cliente_lectura.get(url + "&formato=xlsx")
    hoja = load_workbook(io.BytesIO(excel.content)).active
    assert hoja.max_row >= 2
    pdf = cliente_lectura.get(url + "&formato=pdf")
    assert pdf.content.startswith(b"%PDF")


@pytest.mark.django_db
def test_costo_muestra_porcentaje_al_lado(cliente_lectura, dos_nominas):
    url = reverse("nomina:reportes") + "?reporte=costo&rango=anio&anio=2026&agrupar=departamento"
    html = cliente_lectura.get(url).content.decode()
    assert "Cocina" in html and "6.20%" in html
    hoja = load_workbook(io.BytesIO(cliente_lectura.get(url + "&formato=xlsx").content)).active
    encabezados = [c.value for c in hoja[1]]
    i = encabezados.index("Seguro Social patronal")
    assert encabezados[i + 1] == "% Seguro Social patronal"
