"""Utilidades comunes para importar y exportar hojas Excel/CSV."""

import csv
import io
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


@dataclass
class ErrorFila:
    fila: int
    columna: str
    mensaje: str


def normalizar(texto) -> str:
    """Encabezados y valores sin acentos, en minúscula y con guiones bajos."""
    texto = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode("ascii")
    return texto.strip().lower().replace(" ", "_").replace("-", "_")


def _texto_celda(valor, columna: str, relleno_ceros: dict) -> str:
    if valor is None:
        return ""
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    if isinstance(valor, date):
        return valor.isoformat()
    if isinstance(valor, bool):
        return "si" if valor else "no"
    if isinstance(valor, int):
        texto = str(valor)
        return texto.zfill(relleno_ceros[columna]) if columna in relleno_ceros else texto
    if isinstance(valor, float):
        if valor.is_integer() and columna in relleno_ceros:
            return str(int(valor)).zfill(relleno_ceros[columna])
        return format(Decimal(str(valor)), "f")
    return str(valor).strip()


def leer_archivo(archivo, hoja_preferida: str, relleno_ceros: dict | None = None) -> list[dict]:
    """
    Lee .xlsx o .csv y devuelve filas {columna_normalizada: texto}.
    `relleno_ceros` repone ceros a la izquierda que Excel pierde en números
    (SSN, rutas bancarias, códigos postales).
    """
    relleno_ceros = relleno_ceros or {}
    if archivo.name.lower().endswith(".xlsx"):
        libro = load_workbook(archivo, read_only=True, data_only=True)
        hoja = libro[hoja_preferida] if hoja_preferida in libro.sheetnames else libro.worksheets[0]
        filas = hoja.iter_rows(values_only=True)
        encabezado = [normalizar(c) for c in next(filas, [])]
        datos = [
            {
                encabezado[i]: _texto_celda(v, encabezado[i], relleno_ceros)
                for i, v in enumerate(valores)
                if i < len(encabezado) and encabezado[i]
            }
            for valores in filas
        ]
        libro.close()
        return datos

    crudo = archivo.read()
    try:
        texto = crudo.decode("utf-8-sig")
    except UnicodeDecodeError:
        texto = crudo.decode("cp1252")
    try:
        dialecto = csv.Sniffer().sniff(texto[:4096], delimiters=",;\t")
    except csv.Error:
        dialecto = csv.excel
    lector = csv.reader(io.StringIO(texto), dialecto)
    encabezado = [normalizar(c) for c in next(lector, [])]
    datos = []
    for valores in lector:
        fila = {encabezado[i]: v.strip() for i, v in enumerate(valores) if i < len(encabezado) and encabezado[i]}
        for columna, largo in relleno_ceros.items():
            if fila.get(columna, "").isdigit():
                fila[columna] = fila[columna].zfill(largo)
        datos.append(fila)
    return datos


def fila_vacia(fila: dict) -> bool:
    return not any(v for v in fila.values())


def a_booleano(texto: str) -> bool | None:
    t = normalizar(texto)
    if t in ("", "no", "n", "0", "false", "falso"):
        return False
    if t in ("si", "s", "1", "x", "true", "cierto", "yes", "y"):
        return True
    return None


def opcion(texto: str, choices) -> str:
    """Acepta el código o la etiqueta de una opción (sin importar mayúsculas ni acentos)."""
    t = normalizar(texto)
    for valor, etiqueta in choices:
        if t in (normalizar(valor), normalizar(etiqueta)):
            return valor
    return texto


def celda_segura(valor):
    """Evita que un texto se interprete como fórmula al abrir el Excel exportado."""
    if isinstance(valor, str) and valor[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + valor
    return valor


def plantilla_xlsx(nombre_hoja: str, columnas: list, columnas_texto: tuple, notas: list[str]) -> bytes:
    """
    Plantilla con encabezados (requeridos en amarillo) y una hoja de instrucciones.
    `columnas` = [(columna, requerida, descripción)].
    """
    libro = Workbook()
    hoja = libro.active
    hoja.title = nombre_hoja
    negrita = Font(bold=True)
    amarillo = PatternFill("solid", fgColor="FDE68A")
    nombres = [c[0] for c in columnas]
    for i, (columna, requerida, _desc) in enumerate(columnas, start=1):
        celda = hoja.cell(row=1, column=i, value=columna)
        celda.font = negrita
        if requerida:
            celda.fill = amarillo
        hoja.column_dimensions[celda.column_letter].width = max(14, len(columna) + 2)
    # Formato texto para no perder ceros a la izquierda.
    for columna in columnas_texto:
        letra = hoja.cell(row=1, column=nombres.index(columna) + 1).column_letter
        for fila in range(2, 502):
            hoja[f"{letra}{fila}"].number_format = "@"

    ayuda = libro.create_sheet("Instrucciones")
    ayuda.append(["Columna", "Requerida", "Valores / formato"])
    for celda in ayuda[1]:
        celda.font = negrita
    for columna, requerida, desc in columnas:
        ayuda.append([columna, "Sí" if requerida else "", desc])
    ayuda.append([])
    for nota in notas:
        ayuda.append([nota])
    ayuda.column_dimensions["A"].width = 38
    ayuda.column_dimensions["C"].width = 60

    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def libro_xlsx(nombre_hoja: str, encabezados: list[str], filas: list[list], totales: list | None = None) -> bytes:
    """Exporta una tabla simple a Excel, protegiendo contra inyección de fórmulas."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = nombre_hoja
    hoja.append(encabezados)
    for celda in hoja[1]:
        celda.font = Font(bold=True)
    for fila in filas:
        hoja.append([celda_segura(v) for v in fila])
    if totales:
        hoja.append([celda_segura(v) for v in totales])
        for celda in hoja[hoja.max_row]:
            celda.font = Font(bold=True)
    for i, encabezado in enumerate(encabezados, start=1):
        hoja.column_dimensions[hoja.cell(row=1, column=i).column_letter].width = max(12, len(encabezado) + 4)
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()
