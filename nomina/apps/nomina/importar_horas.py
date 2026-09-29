"""Importar horas e ingresos de un período desde Excel o CSV (p. ej. exportado del reloj ponchador)."""

import io
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

from django.db import transaction
from openpyxl import load_workbook

from apps.core import hojas
from apps.parametros.models import ConceptoIngreso

from .models import EntradaIngreso

HORAS = [
    ("horas_regulares", "Horas regulares"),
    ("horas_extra_diarias", "Horas extra por exceso de 8 al día"),
    ("horas_extra_semanales", "Horas extra por exceso de 40 a la semana"),
    ("horas_septimo_dia", "Horas del séptimo día / día de descanso"),
    ("horas_periodo_alimentos", "Horas trabajadas en el período de tomar alimentos"),
    ("horas_vacaciones", "Horas de vacaciones pagadas"),
    ("horas_enfermedad", "Horas de licencia por enfermedad pagadas"),
]
INGRESOS = [("propinas", "Propinas ($)"), ("comisiones", "Comisiones ($)"), ("bono", "Bono ($)")]
COLUMNAS = ([("numero_empleado", True, "Número de empleado (tal como está en el sistema)"),
             ("empleado", False, "Nombre (sólo como referencia; no se usa)")]
            + [(c, False, d + ". Vacío = 0.") for c, d in HORAS + INGRESOS])


@dataclass
class Error:
    fila: int
    columna: str
    mensaje: str


@dataclass
class Resultado:
    errores: list = field(default_factory=list)
    total_filas: int = 0
    aplicadas: int = 0
    columnas_desconocidas: list = field(default_factory=list)
    columnas_usadas: list = field(default_factory=list)
    vista: list = field(default_factory=list)  # [(entrada, {columna: valor})]


def plantilla(periodo) -> bytes:
    """Plantilla con los empleados del período ya listados."""
    contenido = hojas.plantilla_xlsx(
        "Horas", COLUMNAS, ("numero_empleado",),
        [f"Período {periodo.fecha_inicio:%m/%d/%Y} – {periodo.fecha_fin:%m/%d/%Y}, pago {periodo.fecha_pago:%m/%d/%Y}.",
         "Las columnas que estén en el archivo reemplazan lo entrado para los empleados que aparecen; una celda "
         "vacía es cero. Las columnas que no estén no se tocan. Los empleados que no aparecen no cambian.",
         "Puede borrar las columnas que no use."],
    )
    libro = load_workbook(io.BytesIO(contenido))
    hoja = libro["Horas"]
    # La plantilla ya da formato a las filas vacías; se escribe por posición (append iría al final).
    entradas = periodo.entradas.select_related("empleado").order_by("empleado__apellido_paterno", "empleado__nombre")
    for fila, entrada in enumerate(entradas, start=2):
        hoja.cell(row=fila, column=1, value=entrada.empleado.numero_empleado)
        hoja.cell(row=fila, column=2, value=hojas.celda_segura(entrada.empleado.nombre_completo))
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def _valor(texto):
    texto = (texto or "").strip().replace(",", "").replace("$", "")
    if texto == "":
        return Decimal("0")
    valor = Decimal(texto)
    if valor < 0 or valor > Decimal("100000"):
        raise InvalidOperation
    return valor.quantize(Decimal("0.01"))


def procesar(archivo, periodo, solo_validar=True) -> Resultado:
    resultado = Resultado()
    try:
        filas = hojas.leer_archivo(archivo, "Horas")
    except hojas.ErrorArchivo as e:
        resultado.errores.append(Error(0, "", str(e)))
        return resultado
    except Exception:  # archivo dañado o de otro formato
        resultado.errores.append(Error(0, "", "No se pudo leer el archivo. ¿Es un .xlsx o .csv válido?"))
        return resultado
    conocidas = {c for c, _, _ in COLUMNAS}
    presentes = set().union(*(f.keys() for f in filas)) if filas else set()
    resultado.columnas_desconocidas = sorted(presentes - conocidas)
    usadas = [c for c, _ in HORAS + INGRESOS if c in presentes]
    resultado.columnas_usadas = usadas
    if "numero_empleado" not in presentes:
        resultado.errores.append(Error(1, "numero_empleado", "Falta la columna numero_empleado."))
        return resultado
    entradas = {e.empleado.numero_empleado: e for e in periodo.entradas.select_related("empleado")}
    vistos = {}
    for n, fila in enumerate(filas, start=2):
        if hojas.fila_vacia(fila):
            continue
        resultado.total_filas += 1
        numero = (fila.get("numero_empleado") or "").strip()
        entrada = entradas.get(numero)
        if entrada is None:
            resultado.errores.append(Error(n, "numero_empleado", f"El empleado «{numero}» no está en este período."))
            continue
        if numero in vistos:
            resultado.errores.append(Error(n, "numero_empleado", f"El empleado {numero} se repite (fila {vistos[numero]})."))
            continue
        vistos[numero] = n
        valores = {}
        for columna in usadas:
            try:
                valores[columna] = _valor(fila.get(columna))
            except (InvalidOperation, ValueError):
                resultado.errores.append(Error(n, columna, f"«{fila.get(columna)}» no es una cantidad válida."))
        resultado.vista.append((entrada, valores))
    if resultado.errores or solo_validar:
        return resultado
    conceptos = {c.codigo: c for c in ConceptoIngreso.objects.filter(codigo__in=[c for c, _ in INGRESOS])}
    with transaction.atomic():
        for entrada, valores in resultado.vista:
            for columna, valor in valores.items():
                if columna in dict(HORAS):
                    setattr(entrada, columna, valor)
            entrada.incluir = True
            entrada.save()
            for codigo in (c for c in valores if c in conceptos):
                entrada.ingresos.filter(concepto=conceptos[codigo]).delete()
                if valores[codigo]:
                    EntradaIngreso.objects.create(entrada=entrada, concepto=conceptos[codigo], monto=valores[codigo],
                                                  descripcion="Importado")
            resultado.aplicadas += 1
    return resultado
