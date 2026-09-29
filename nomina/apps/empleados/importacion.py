"""
Importación masiva de empleados desde Excel (.xlsx) o CSV.

Reglas:
- Cada fila se valida con el mismo formulario que la pantalla (mismas reglas).
- "Todo o nada": si alguna fila tiene errores no se importa ninguna.
- Solo crea empleados nuevos; un número de empleado existente es un error.
- Departamentos y clasificaciones CFSE que no existan se crean.
"""

import csv
import io
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal

from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from apps.companias.models import ClasificacionCFSE, Departamento

from .forms import EmpleadoForm
from .models import Empleado

MAX_FILAS = 2000

# (columna del archivo, campo del formulario, requerido, descripción/valores)
COLUMNAS = [
    ("numero_empleado", "numero_empleado", True, "Número único dentro de la compañía"),
    ("nombre", "nombre", True, ""),
    ("segundo_nombre", "segundo_nombre", False, ""),
    ("apellido_paterno", "apellido_paterno", True, ""),
    ("apellido_materno", "apellido_materno", False, ""),
    ("ssn", "ssn_nuevo", True, "9 dígitos, con o sin guiones"),
    ("fecha_nacimiento", "fecha_nacimiento", False, "AAAA-MM-DD o MM/DD/AAAA"),
    ("fecha_empleo", "fecha_empleo", True, "AAAA-MM-DD o MM/DD/AAAA"),
    ("regimen_laboral", "regimen_laboral", False, "anterior | ley4 | vacío = automático por fecha"),
    ("direccion", "direccion_linea1", False, ""),
    ("direccion_2", "direccion_linea2", False, ""),
    ("ciudad", "ciudad", False, ""),
    ("estado", "estado", False, "PR por defecto"),
    ("codigo_postal", "codigo_postal", False, ""),
    ("telefono", "telefono", False, ""),
    ("email", "email", False, ""),
    ("departamento", None, False, "Nombre; se crea si no existe"),
    ("clasificacion_cfse", None, False, "Código; se crea si no existe"),
    ("tipo_pago", "tipo_pago", True, "hora | salario | exento"),
    ("tarifa", "tarifa", True, "Por hora, o salario por período"),
    ("horas_regulares_periodo", "horas_regulares_periodo", False, "Solo asalariados"),
    ("aplica_choferil", "aplica_choferil", False, "si | no"),
    ("recibe_propinas", "recibe_propinas", False, "si | no"),
    ("r4_estado_civil", "r4_estado_civil", False, "soltero | casado | casado_separado | jefe_familia"),
    ("r4_exencion_personal", "r4_exencion_personal", False, "completa | mitad | ninguna"),
    ("r4_dependientes", "r4_dependientes", False, "Número"),
    ("r4_dependientes_custodia_compartida", "r4_dependientes_custodia_compartida", False, "Número"),
    ("r4_veterano", "r4_veterano", False, "si | no"),
    ("r4_concesion_deducciones", "r4_concesion_deducciones", False, "$ anual"),
    ("r4_retencion_adicional", "r4_retencion_adicional", False, "$ por período"),
    ("w4_aplica", "w4_aplica", False, "si | no"),
    ("w4_estado_civil", "w4_estado_civil", False, "single | married | head"),
    ("w4_paso2", "w4_paso2", False, "si | no"),
    ("w4_dependientes", "w4_dependientes", False, "$"),
    ("w4_otros_ingresos", "w4_otros_ingresos", False, "$"),
    ("w4_deducciones", "w4_deducciones", False, "$"),
    ("w4_retencion_adicional", "w4_retencion_adicional", False, "$ por período"),
    ("deposito_directo", "deposito_directo", False, "si | no"),
    ("banco_nombre", "banco_nombre", False, ""),
    ("banco_ruta", "banco_ruta", False, "9 dígitos"),
    ("banco_cuenta", "cuenta_nueva", False, ""),
    ("banco_tipo_cuenta", "banco_tipo_cuenta", False, "cheques | ahorros"),
    ("notas", "notas", False, ""),
]
BOOLEANOS = {
    "aplica_choferil", "recibe_propinas", "r4_veterano", "w4_aplica", "w4_paso2", "deposito_directo",
}
VALORES_POR_DEFECTO = {
    "estado": "PR",
    "r4_estado_civil": Empleado.EstadoCivilPR.SOLTERO,
    "r4_exencion_personal": Empleado.ExencionPersonal.COMPLETA,
    "r4_dependientes": "0",
    "r4_dependientes_custodia_compartida": "0",
    "r4_concesion_deducciones": "0",
    "r4_retencion_adicional": "0",
    "w4_estado_civil": Empleado.EstadoCivilW4.SOLTERO,
    "w4_dependientes": "0",
    "w4_otros_ingresos": "0",
    "w4_deducciones": "0",
    "w4_retencion_adicional": "0",
    "banco_tipo_cuenta": Empleado.TipoCuenta.CHEQUES,
}
COLUMNAS_RELLENO_CEROS = {"ssn": 9, "banco_ruta": 9, "codigo_postal": 5}


def _normalizar(texto) -> str:
    texto = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode("ascii")
    return texto.strip().lower().replace(" ", "_").replace("-", "_")


@dataclass
class ErrorFila:
    fila: int
    columna: str
    mensaje: str


@dataclass
class Resultado:
    total_filas: int = 0
    errores: list = field(default_factory=list)
    departamentos_nuevos: set = field(default_factory=set)
    cfse_nuevos: set = field(default_factory=set)
    creados: list = field(default_factory=list)
    columnas_desconocidas: list = field(default_factory=list)

    @property
    def ok(self):
        return not self.errores and self.total_filas > 0


def _texto_celda(valor, columna: str) -> str:
    if valor is None:
        return ""
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    if isinstance(valor, date):
        return valor.isoformat()
    if isinstance(valor, bool):
        return "si" if valor else "no"
    if isinstance(valor, int):
        texto = str(valor)
        if columna in COLUMNAS_RELLENO_CEROS:
            texto = texto.zfill(COLUMNAS_RELLENO_CEROS[columna])
        return texto
    if isinstance(valor, float):
        if valor.is_integer() and columna in COLUMNAS_RELLENO_CEROS:
            return str(int(valor)).zfill(COLUMNAS_RELLENO_CEROS[columna])
        return format(Decimal(str(valor)), "f")
    return str(valor).strip()


def leer_archivo(archivo) -> list[dict]:
    """Devuelve una lista de filas {columna_normalizada: texto}."""
    nombre = archivo.name.lower()
    if nombre.endswith(".xlsx"):
        libro = load_workbook(archivo, read_only=True, data_only=True)
        hoja = libro["Empleados"] if "Empleados" in libro.sheetnames else libro.worksheets[0]
        filas = hoja.iter_rows(values_only=True)
        encabezado = [_normalizar(c) for c in next(filas, [])]
        datos = []
        for valores in filas:
            fila = {
                encabezado[i]: _texto_celda(v, encabezado[i])
                for i, v in enumerate(valores)
                if i < len(encabezado) and encabezado[i]
            }
            datos.append(fila)
        libro.close()
        return datos

    crudo = archivo.read()
    try:
        texto = crudo.decode("utf-8-sig")
    except UnicodeDecodeError:
        texto = crudo.decode("cp1252")
    try:
        dialecto = csv.Sniffer().sniff(texto[:4096], delimiters=",;\t")
    except csv.Error:
        dialecto = csv.excel
    lector = csv.reader(io.StringIO(texto), dialecto)
    encabezado = [_normalizar(c) for c in next(lector, [])]
    return [
        {encabezado[i]: v.strip() for i, v in enumerate(valores) if i < len(encabezado) and encabezado[i]}
        for valores in lector
    ]


def _a_booleano(texto: str) -> bool | None:
    t = _normalizar(texto)
    if t in ("", "no", "n", "0", "false", "falso"):
        return False
    if t in ("si", "s", "1", "x", "true", "cierto", "yes", "y"):
        return True
    return None


def _opcion(texto: str, choices) -> str:
    """Acepta el código o la etiqueta (sin importar mayúsculas ni acentos)."""
    t = _normalizar(texto)
    for valor, etiqueta in choices:
        if t in (_normalizar(valor), _normalizar(etiqueta)):
            return valor
    return texto


def _fila_vacia(fila: dict) -> bool:
    return not any(v for v in fila.values())


def procesar(archivo, compania, usuario, solo_validar=True) -> Resultado:
    resultado = Resultado()
    try:
        filas = leer_archivo(archivo)
    except Exception:  # archivo dañado o formato inesperado
        resultado.errores.append(ErrorFila(0, "", "No se pudo leer el archivo. Verifique que sea .xlsx o .csv válido."))
        return resultado

    conocidas = {c[0] for c in COLUMNAS}
    if filas:
        resultado.columnas_desconocidas = sorted(set(filas[0]) - conocidas)
        faltan = [c[0] for c in COLUMNAS if c[2] and c[0] not in filas[0]]
        if faltan:
            resultado.errores.append(ErrorFila(1, ", ".join(faltan), "Faltan columnas requeridas en el encabezado."))
            return resultado

    campos_modelo = {f.name: f for f in Empleado._meta.fields}
    deptos = {d.nombre.lower(): d for d in Departamento.objects.filter(compania=compania)}
    cfse = {c.codigo.lower(): c for c in ClasificacionCFSE.objects.filter(compania=compania)}
    ssn_vistos, numeros_vistos = {}, {}
    validos = []

    for indice, fila in enumerate(filas, start=2):  # fila 1 = encabezado
        if _fila_vacia(fila):
            continue
        resultado.total_filas += 1
        if resultado.total_filas > MAX_FILAS:
            resultado.errores.append(ErrorFila(indice, "", f"Máximo {MAX_FILAS} empleados por archivo."))
            break

        datos = {}
        for columna, campo, _req, _desc in COLUMNAS:
            if campo is None:
                continue
            valor = fila.get(columna, "")
            if not valor and campo in VALORES_POR_DEFECTO:
                valor = VALORES_POR_DEFECTO[campo]
            if campo in BOOLEANOS:
                booleano = _a_booleano(valor)
                if booleano is None:
                    resultado.errores.append(ErrorFila(indice, columna, f"Valor '{valor}' no es si/no."))
                    continue
                if booleano:
                    datos[campo] = "on"
                continue
            modelo = campos_modelo.get(campo)
            if valor and modelo is not None and modelo.choices:
                valor = _opcion(valor, modelo.choices)
            datos[campo] = valor

        form = EmpleadoForm(datos, compania=compania)
        if not form.is_valid():
            columna_de = {c[1]: c[0] for c in COLUMNAS if c[1]}
            for campo, mensajes in form.errors.items():
                columna = columna_de.get(campo, "") if campo != "__all__" else ""
                for mensaje in mensajes:
                    resultado.errores.append(ErrorFila(indice, columna, mensaje))
            continue

        ssn = form.cleaned_data["ssn_nuevo"]
        if ssn in ssn_vistos:
            resultado.errores.append(ErrorFila(indice, "ssn", f"SSN repetido en el archivo (fila {ssn_vistos[ssn]})."))
            continue
        ssn_vistos[ssn] = indice
        numero = form.cleaned_data["numero_empleado"]
        if numero in numeros_vistos:
            resultado.errores.append(
                ErrorFila(indice, "numero_empleado", f"Número repetido en el archivo (fila {numeros_vistos[numero]}).")
            )
            continue
        numeros_vistos[numero] = indice

        depto = fila.get("departamento", "").strip()
        if depto and depto.lower() not in deptos:
            resultado.departamentos_nuevos.add(depto)
        codigo_cfse = fila.get("clasificacion_cfse", "").strip()
        if codigo_cfse and codigo_cfse.lower() not in cfse:
            resultado.cfse_nuevos.add(codigo_cfse)
        validos.append((form, depto, codigo_cfse))

    if resultado.errores or solo_validar or not validos:
        return resultado

    with transaction.atomic():
        for form, depto, codigo_cfse in validos:
            empleado = form.save(commit=False)
            empleado.creado_por = usuario
            if depto:
                clave = depto.lower()
                if clave not in deptos:
                    deptos[clave] = Departamento.objects.create(compania=compania, nombre=depto)
                empleado.departamento = deptos[clave]
            if codigo_cfse:
                clave = codigo_cfse.lower()
                if clave not in cfse:
                    cfse[clave] = ClasificacionCFSE.objects.create(compania=compania, codigo=codigo_cfse)
                empleado.clasificacion_cfse = cfse[clave]
            empleado.save()
            resultado.creados.append(empleado)
    return resultado


def plantilla_xlsx() -> bytes:
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Empleados"
    negrita = Font(bold=True)
    requerido = PatternFill("solid", fgColor="FDE68A")
    for i, (columna, _campo, req, _desc) in enumerate(COLUMNAS, start=1):
        celda = hoja.cell(row=1, column=i, value=columna)
        celda.font = negrita
        if req:
            celda.fill = requerido
        hoja.column_dimensions[celda.column_letter].width = max(14, len(columna) + 2)
    # Formato texto para no perder ceros a la izquierda.
    for columna in ("ssn", "banco_ruta", "banco_cuenta", "codigo_postal", "numero_empleado"):
        letra = hoja.cell(row=1, column=[c[0] for c in COLUMNAS].index(columna) + 1).column_letter
        for fila in range(2, 502):
            hoja[f"{letra}{fila}"].number_format = "@"

    ayuda = libro.create_sheet("Instrucciones")
    ayuda.append(["Columna", "Requerida", "Valores / formato"])
    for celda in ayuda[1]:
        celda.font = negrita
    for columna, _campo, req, desc in COLUMNAS:
        ayuda.append([columna, "Sí" if req else "", desc])
    ayuda.append([])
    ayuda.append(["Las columnas en amarillo son requeridas. Un empleado por fila."])
    ayuda.append(["Si alguna fila tiene errores, no se importa ninguna. Corrija y vuelva a subir."])
    ayuda.column_dimensions["A"].width = 38
    ayuda.column_dimensions["C"].width = 60

    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()
